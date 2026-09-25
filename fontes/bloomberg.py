"""Ponte Bloomberg: gera o template Excel com fórmulas BDP e lê os valores calculados.

Fluxo (na máquina com terminal):
  1. `python coletar.py --ponte`  -> cria consenso/ponte_bloomberg.xlsx (fórmulas BDP)
  2. Abrir o arquivo no Excel com o add-in Bloomberg, esperar calcular, salvar.
  3. `python coletar.py --consenso-bloomberg` -> lê os valores e acrescenta em consenso/consenso.csv
     (uma linha por ticker/ano/data; o histórico de revisões nasce daí).

Licença: dado Bloomberg fica na estação/na firma. Repositório privado, dashboard local.
"""
from __future__ import annotations

from datetime import date

import config

PONTE = config.CONSENSO / "ponte_bloomberg.xlsx"


def gera_ponte() -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "consenso"
    cab = ["ticker", "ano", "bbg_ticker"] + list(config.BLOOMBERG_CAMPOS) + ["preco"]
    ws.append(cab)
    r = 2
    for tk in config.UNIVERSO:
        for ano in config.ANOS_FISCAIS:
            bbg = f"{tk} BZ Equity"
            ws.cell(r, 1, tk)
            ws.cell(r, 2, ano)
            ws.cell(r, 3, bbg)
            for j, campo in enumerate(config.BLOOMBERG_CAMPOS.values(), start=4):
                if campo in ("BEST_TARGET_PRICE", "TOT_ANALYST_REC"):
                    ws.cell(r, j, f'=BDP($C{r},"{campo}")')
                else:
                    # override de período fiscal: 2026 -> "2026FY"
                    ws.cell(r, j, f'=BDP($C{r},"{campo}","BEST_FPERIOD_OVERRIDE","{ano}FY")')
            ws.cell(r, len(cab), f'=BDP($C{r},"PX_LAST")')
            r += 1
    ws.append([])
    ws.append(["Obs: receita/EBITDA/lucro vêm em MILHÕES de BRL (conferir BEST_* currency/scale no terminal)."])
    # aba macro: P/L à frente do Ibov (para Ibov = Lucro/(Ke-g)) e CDS Brasil 5a
    wm = wb.create_sheet("macro")
    wm.append(["chave", "bbg_ticker", "campo", "valor"])
    for chave, tk, campo in (("ibov_pe_fwd", "IBOV Index", "BEST_PE_RATIO"), ("ibov_eps_fwd", "IBOV Index", "BEST_EPS"),
                             ("cds_brasil_5a", "CBRZ1U5 CBGN Curncy", "PX_LAST"), ("ibov_px", "IBOV Index", "PX_LAST")):
        r = wm.max_row + 1
        wm.append([chave, tk, campo, f'=BDP($B{r},$C{r})'])
    wb.save(PONTE)
    print(f"  ponte gerada: {PONTE}")


def le_ponte_macro() -> dict:
    """Valores da aba 'macro' da ponte -> data/bloomberg_macro.json (com data)."""
    import json
    import openpyxl
    if not PONTE.exists():
        return {}
    wb = openpyxl.load_workbook(PONTE, data_only=True)
    if "macro" not in wb.sheetnames:
        return {}
    out = {"data": date.today().isoformat()}
    for chave, _, _, valor in wb["macro"].iter_rows(min_row=2, values_only=True):
        if chave and isinstance(valor, (int, float)):
            out[chave] = float(valor)
    (config.DATA / "bloomberg_macro.json").write_text(json.dumps(out), encoding="utf-8")
    return out


def le_ponte() -> list[dict]:
    """Lê valores calculados (data_only) da ponte. Linhas sem número são ignoradas."""
    import openpyxl
    if not PONTE.exists():
        print("  ponte não existe; rode --ponte primeiro")
        return []
    wb = openpyxl.load_workbook(PONTE, data_only=True)
    ws = wb["consenso"]
    cab = [c.value for c in ws[1]]
    hoje = date.today().isoformat()
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not isinstance(row[1], (int, float)):
            continue
        d = dict(zip(cab, row))
        vals = {k: (float(d[k]) if isinstance(d.get(k), (int, float)) else None) for k in config.BLOOMBERG_CAMPOS}
        if all(v is None for v in vals.values()):
            continue  # fórmula não calculou (abrir no Excel com add-in)
        out.append({
            "ticker": d["ticker"], "data": hoje, "ano": int(d["ano"]),
            "receita": vals["receita"], "ebitda": vals["ebitda"], "lucro": vals["lucro"],
            "eps": vals["eps"], "target": vals["target"],
            "n_analistas": int(vals["n_analistas"]) if vals["n_analistas"] else None,
            "fonte": "bloomberg",
        })
    return out
